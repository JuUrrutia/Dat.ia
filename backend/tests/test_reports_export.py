"""
Integridad de los documentos ejecutivos (Excel + PDF).

Regla del modulo: ningun valor por defecto puede afirmar un resultado que no se
verifico, y ningun dato de usuario o del LLM puede viajar crudo a un formato que
lo interpreta (parser XML del PDF, motor de formulas del Excel).

Dos capas, en dos clases:

- `TestExportHttpContract`: lo que el endpoint RESPONDE (status, content-type,
  content-disposition, 401 sin token, 404/403, 409 con snapshot no exportable).
- `TestExportContent`: lo que el documento DICE una vez renderizado.
"""
import importlib.util
import io
import json
import re
import unittest
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app
from app.core.database import SessionLocal
from app.db.init_db import init_db
from app.modules.auth.models import User, UserSession, Role
from app.modules.telemetry_audit.models import AuditLog
from app.core.security import create_access_token, get_password_hash
from app.modules.admin_catalog.schemas import (
    ExecutiveReportData,
    KPICardData,
    ReportExportData,
    TraceabilityAuditData,
)
from app.modules.reports.excel_exporter import ExcelExporter
from app.modules.reports.pdf_exporter import PDFExporter

HAS_PYMUPDF = importlib.util.find_spec("fitz") is not None

# Todo lo que un LLM o un usuario podrian escribir de verdad: markup, ampersand,
# y prefijos que Excel/Sheets convierten en formula.
HOSTILE_QUESTION = 'x <img src=y> & z'
FORMULA_QUESTION = '=HYPERLINK("http://evil","click")'
HOSTILE_SQL = '=1+1'
HOSTILE_KPI_TITLE = '<b>ALTAS</b> & mas'

# Titulos de seccion del PDF, para contar la numeracion sin contar las
# recomendaciones numeradas (que tambien arrancan con "N. ").
SECTION_RE = re.compile(
    r"^(\d+)\. (Diagn|Hallazgos|Recomendaciones|Visualizaci|Trazabilidad)", re.M
)

# Snapshot completo (una consulta de ventas ya resuelta): lo que el servidor
# persiste en `AuditLog.result_snapshot` y de donde sale el documento.
SAMPLE_SNAPSHOT_DATA = {
    "question": "¿Cuáles fueron los 5 productos más vendidos del trimestre?",
    "summary_text": "Los 5 productos más vendidos totalizaron 12,500 unidades vendidas.",
    "executive_report": {
        "overview": "El desempeño del trimestre muestra un crecimiento sólido en la línea de periféricos.",
        "key_findings": [
            "El producto 'Teclado Mecánico RGB' lidera el volumen con 4,200 unidades.",
            "El ticket promedio aumentó un 12% interanual.",
            "El canal online representa el 68% de las conversiones totales."
        ],
        "recommendations": [
            "Aumentar el stock de seguridad para el producto líder.",
            "Lanzar campaña de cross-selling con monitores gamer.",
            "Revisar márgenes de distribución en el canal retail."
        ],
        "risk_level": "MEDIO",
        "business_impact": "Generación proyectada de $45,000 USD en ingresos adicionales."
    },
    "kpis": [
        {"title": "Ventas Totales", "value": "$145,200", "subtitle": "+15% vs Q anterior"},
        {"title": "Unidades", "value": "12,500", "subtitle": "Top 5 productos"},
        {"title": "Margen Bruto", "value": "34.5%", "subtitle": "Estable"}
    ],
    "data_columns": ["id_producto", "nombre_producto", "unidades_vendidas", "total_ingresos"],
    "data_rows": [
        {"id_producto": 101, "nombre_producto": "Teclado Mecánico RGB", "unidades_vendidas": 4200, "total_ingresos": 42000.0},
        {"id_producto": 102, "nombre_producto": "Mouse Gamer Pro", "unidades_vendidas": 3100, "total_ingresos": 24800.0},
        {"id_producto": 103, "nombre_producto": "Monitor 27 UHD", "unidades_vendidas": 2200, "total_ingresos": 55000.0},
        {"id_producto": 104, "nombre_producto": "Auriculares Wireless", "unidades_vendidas": 1800, "total_ingresos": 14400.0},
        {"id_producto": 105, "nombre_producto": "Mousepad XL", "unidades_vendidas": 1200, "total_ingresos": 9000.0}
    ],
    "traceability": {
        "sql_executed": "SELECT id_producto, nombre_producto, unidades_vendidas, total_ingresos FROM ventas ORDER BY unidades_vendidas DESC LIMIT 5",
        "execution_time_ms": 14,
        "rows_returned": 5,
        "validation_status": "APROBADO",
        "schema_tables_used": ["ventas"]
    },
    "target_database": "demo_corporativa.db"
}

# Snapshot minimo para la auditoria: `traceability.rows_returned` (5000) no
# concuerda con las 5 filas del snapshot a proposito, y los tests de auditoria
# justamente comparan esas dos cifras.
SNAPSHOT = {
    "question": "¿Cuáles fueron los 5 productos más vendidos?",
    "summary_text": "Totalizaron 12,500 unidades.",
    "kpis": [{"title": "Ventas", "value": "$145,200"}],
    "data_columns": ["id", "nombre"],
    "data_rows": [
        {"id": 1, "nombre": "Teclado"},
        {"id": 2, "nombre": "Mouse"},
        {"id": 3, "nombre": "Monitor"},
        {"id": 4, "nombre": "Auriculares"},
        {"id": 5, "nombre": "Mousepad"},
    ],
    "traceability": {
        "sql_executed": "SELECT id, nombre FROM ventas",
        "execution_time_ms": 14,
        "rows_returned": 5000,
        "validation_status": "APROBADO",
    },
    "target_database": "postgres://cliente",
}


def _section_numbers(text: str):
    return [int(m.group(1)) for m in SECTION_RE.finditer(text)]


def _data(**overrides) -> ReportExportData:
    base = dict(
        question=HOSTILE_QUESTION,
        summary_text="ok",
        executive_report=ExecutiveReportData(
            overview=f"{HOSTILE_KPI_TITLE} — riesgo & retorno",
            key_findings=["<b>MALAS</b> & bajas"],
            recommendations=[f"{i}. <i>bien</i> & mal" for i in range(1, 4)],
            risk_level="MEDIO & ALTO",
            business_impact="Impacto <b>fuerte</b> & medible",
        ),
        kpis=[KPICardData(title=f"KPI {i}", value="=1+1", subtitle=f"<i>sub</i> {i} & mas") for i in range(7)],
        data_columns=["a"],
        data_rows=[{"a": "=1+1"}, {"a": "normal"}],
        traceability=TraceabilityAuditData(
            sql_executed=HOSTILE_SQL,
            execution_time_ms=14,
            rows_returned=2,
            validation_status="APROBADO",
        ),
        target_database=None,
        custom_notes="nota <img src=y> & z",
    )
    base.update(overrides)
    return ReportExportData(**base)


def _cells(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    return [(ws.title, c.coordinate, c) for ws in wb.worksheets for row in ws.iter_rows() for c in row]


def _pdf_text(content: bytes) -> str:
    import fitz  # pymupdf, no es dependencia dura del proyecto
    return fitz.open(stream=content, filetype="pdf")[0].get_text()


class TestExportHttpContract(unittest.TestCase):
    """El contrato HTTP del export: el archivo llega, con el tipo y el nombre correctos,
    y solo a quien puede verlo."""

    @classmethod
    def setUpClass(cls):
        db = SessionLocal()
        init_db(db)
        db.close()

    def setUp(self):
        self.client = TestClient(app)
        self.db = SessionLocal()

        self.user = self.db.query(User).filter(User.username == "admin").first()
        self.jti = str(uuid.uuid4())
        self.token = create_access_token(subject=self.user.id, jti=self.jti)

        session = UserSession(
            user_id=self.user.id,
            jti=self.jti,
            is_revoked=False
        )
        self.db.add(session)
        self.db.commit()

        self.headers = {"Authorization": f"Bearer {self.token}"}

        # AuditLog con result_snapshot para el usuario admin
        self.audit_log = AuditLog(
            user_id=self.user.id,
            username=self.user.username,
            user_role="Administrador",
            question_prompt=SAMPLE_SNAPSHOT_DATA["question"],
            sql_generated=SAMPLE_SNAPSHOT_DATA["traceability"]["sql_executed"],
            validation_status="APROBADO",
            target_database="demo_corporativa.db",
            execution_time_ms=14,
            rows_returned=5,
            result_snapshot=json.dumps(SAMPLE_SNAPSHOT_DATA)
        )
        self.db.add(self.audit_log)
        self.db.commit()
        self.db.refresh(self.audit_log)

    def tearDown(self):
        if hasattr(self, "audit_log") and self.audit_log.id:
            self.db.query(AuditLog).filter(AuditLog.id == self.audit_log.id).delete()
            self.db.commit()
        self.db.close()

    def test_export_pdf_success(self):
        """Genera el PDF desde el snapshot persistido en auditoria, por audit_log_id."""
        chart_base64 = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        response = self.client.post(
            "/api/v1/reports/export/pdf",
            json={"audit_log_id": self.audit_log.id, "chart_image_base64": chart_base64},
            headers=self.headers
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("application/pdf", response.headers["content-type"])
        self.assertIn("attachment; filename=informe_ejecutivo_dat.ia_", response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"%PDF-"))

        # Verifica que la entrada de auditoria del export quedo registrada
        log = self.db.query(AuditLog).filter(AuditLog.validation_status == "EXPORTADO_PDF").order_by(AuditLog.id.desc()).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.username, "admin")
        self.assertEqual(log.question_prompt, SAMPLE_SNAPSHOT_DATA["question"])

    def test_export_excel_success(self):
        """Genera el Excel desde el snapshot persistido en auditoria, por audit_log_id."""
        response = self.client.post(
            "/api/v1/reports/export/excel",
            json={"audit_log_id": self.audit_log.id},
            headers=self.headers
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response.headers["content-type"])
        self.assertIn("attachment; filename=datos_dat.ia_", response.headers["content-disposition"])
        # Magic header de ZIP / OpenXML: PK\x03\x04
        self.assertTrue(response.content.startswith(b"PK\x03\x04"))

        # Verifica que la entrada de auditoria del export quedo registrada
        log = self.db.query(AuditLog).filter(AuditLog.validation_status == "EXPORTADO_EXCEL").order_by(AuditLog.id.desc()).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.username, "admin")
        self.assertEqual(log.rows_returned, 5)

    def test_export_endpoints_require_authentication(self):
        """Sin header JWT, los endpoints de export responden 401."""
        resp_pdf = self.client.post("/api/v1/reports/export/pdf", json={"audit_log_id": self.audit_log.id})
        self.assertEqual(resp_pdf.status_code, 401)

        resp_excel = self.client.post("/api/v1/reports/export/excel", json={"audit_log_id": self.audit_log.id})
        self.assertEqual(resp_excel.status_code, 401)

    def test_export_nonexistent_audit_log_returns_404(self):
        """Exportar un audit_log_id inexistente responde 404 Not Found."""
        response = self.client.post(
            "/api/v1/reports/export/pdf",
            json={"audit_log_id": 99999999},
            headers=self.headers
        )
        self.assertEqual(response.status_code, 404)
        self.assertIn("Registro de auditoría no encontrado", response.json().get("detail", ""))

    def test_cannot_export_other_user_audit_log(self):
        """
        Un usuario no-admin no puede exportar el result_snapshot de otro usuario.
        Debe responder 403 Forbidden.
        """
        # 1. User A (no-admin)
        usuario_role = self.db.query(Role).filter(Role.name == "Usuario").first()
        user_a = User(
            username=f"user_a_{uuid.uuid4().hex[:6]}",
            hashed_password=get_password_hash("Password123!"),
            is_admin=False,
            is_active=True,
            role_id=usuario_role.id if usuario_role else None
        )
        # 2. User B (no-admin)
        user_b = User(
            username=f"user_b_{uuid.uuid4().hex[:6]}",
            hashed_password=get_password_hash("Password123!"),
            is_admin=False,
            is_active=True,
            role_id=usuario_role.id if usuario_role else None
        )
        self.db.add(user_a)
        self.db.add(user_b)
        self.db.commit()
        self.db.refresh(user_a)
        self.db.refresh(user_b)

        # 3. AuditLog que pertenece a User A
        log_a = AuditLog(
            user_id=user_a.id,
            username=user_a.username,
            user_role="Usuario",
            question_prompt="Consulta confidencial de User A",
            sql_generated="SELECT 1",
            validation_status="APROBADO",
            target_database="demo_corporativa.db",
            execution_time_ms=5,
            rows_returned=1,
            result_snapshot=json.dumps(SAMPLE_SNAPSHOT_DATA)
        )
        self.db.add(log_a)
        self.db.commit()
        self.db.refresh(log_a)

        # 4. JWT de User B
        jti_b = str(uuid.uuid4())
        token_b = create_access_token(subject=user_b.id, jti=jti_b)
        session_b = UserSession(user_id=user_b.id, jti=jti_b, is_revoked=False)
        self.db.add(session_b)
        self.db.commit()
        headers_b = {"Authorization": f"Bearer {token_b}"}

        try:
            # User B intenta exportar el audit log de User A
            resp_pdf = self.client.post(
                "/api/v1/reports/export/pdf",
                json={"audit_log_id": log_a.id},
                headers=headers_b
            )
            self.assertEqual(resp_pdf.status_code, 403)
            self.assertIn("No tienes permisos para exportar consultas de otro usuario", resp_pdf.json().get("detail", ""))

            resp_excel = self.client.post(
                "/api/v1/reports/export/excel",
                json={"audit_log_id": log_a.id},
                headers=headers_b
            )
            self.assertEqual(resp_excel.status_code, 403)
            self.assertIn("No tienes permisos para exportar consultas de otro usuario", resp_excel.json().get("detail", ""))
        finally:
            self.db.query(AuditLog).filter(AuditLog.id == log_a.id).delete()
            self.db.query(UserSession).filter(UserSession.user_id == user_b.id).delete()
            self.db.query(User).filter(User.id.in_([user_a.id, user_b.id])).delete()
            self.db.commit()


class TestExportContent(unittest.TestCase):
    """El documento renderizado: lo que dice no puede afirmar mas de lo que sabe,
    y lo que el usuario escribio viaja como texto, no como markup."""

    @classmethod
    def setUpClass(cls):
        db = SessionLocal()
        init_db(db)
        db.close()

    def setUp(self):
        self.client = TestClient(app)
        self.db = SessionLocal()
        self.user = self.db.query(User).filter(User.username == "admin").first()
        self.jti = str(uuid.uuid4())
        self.token = create_access_token(subject=self.user.id, jti=self.jti)
        self.db.add(UserSession(user_id=self.user.id, jti=self.jti, is_revoked=False))
        self.headers = {"Authorization": f"Bearer {self.token}"}
        self.created_ids = []
        self.db.commit()

    def tearDown(self):
        for audit_id in self.created_ids:
            self.db.query(AuditLog).filter(AuditLog.id == audit_id).delete()
        self.db.query(UserSession).filter(UserSession.jti == self.jti).delete()
        self.db.commit()
        self.db.close()

    def _log(self, snapshot) -> AuditLog:
        entry = AuditLog(
            user_id=self.user.id,
            username=self.user.username,
            user_role="Administrador",
            question_prompt=SNAPSHOT["question"],
            sql_generated=SNAPSHOT["traceability"]["sql_executed"],
            validation_status="APROBADO",
            target_database="postgres://cliente",
            execution_time_ms=14,
            rows_returned=5000,
            result_snapshot=snapshot,
        )
        self.db.add(entry)
        self.db.commit()
        self.db.refresh(entry)
        self.created_ids.append(entry.id)
        return entry

    def _last_export_log(self, status: str) -> AuditLog:
        self.db.expire_all()
        return (
            self.db.query(AuditLog)
            .filter(AuditLog.validation_status == status)
            .order_by(AuditLog.id.desc())
            .first()
        )

    # --- BUG 2 y BUG 3: lo que no se sabe se dice, no se inventa ----------

    def test_missing_target_database_is_not_named_after_the_demo(self):
        for data in (_data(), _data(target_database=None)):
            for content in (ExcelExporter.generate_excel(data), PDFExporter.generate_pdf(data)):
                self.assertNotIn(b"demo_corporativa.db", content)
                self.assertNotIn(b"SQLite Demo", content)

    def test_missing_target_database_says_not_available(self):
        xlsx = ExcelExporter.generate_excel(_data())
        summary = load_workbook(io.BytesIO(xlsx))["Resumen Ejecutivo"]
        self.assertEqual(summary["B6"].value, "No disponible")

    def test_without_traceability_no_approved_status_is_printed(self):
        for data in (_data(traceability=None), _data(traceability=TraceabilityAuditData(sql_executed=HOSTILE_SQL, execution_time_ms=None, rows_returned=None, validation_status=None))):
            xlsx = ExcelExporter.generate_excel(data)
            pdf = PDFExporter.generate_pdf(data)
            for content in (xlsx, pdf):
                self.assertNotIn(b"APROBADO", content)

    def test_without_traceability_it_states_the_data_is_missing(self):
        data = _data(traceability=None)
        summary = load_workbook(io.BytesIO(ExcelExporter.generate_excel(data)))["Resumen Ejecutivo"]
        trace_line = next(
            c.value for row in summary.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("Estado AST:")
        )
        self.assertIn("Sin registro de validación", trace_line)
        self.assertIn("Filas: No disponible", trace_line)
        self.assertIn("Latencia: No disponible", trace_line)

    def test_summary_text_never_claims_success_it_cannot_back(self):
        data = _data(summary_text=None, executive_report=None)
        self.assertNotIn(b"satisfactoriamente", ExcelExporter.generate_excel(data))
        self.assertNotIn(b"satisfactoriamente", PDFExporter.generate_pdf(data))

    # --- BUG 6: el texto del usuario/LLM viaja como texto, no como markup --

    @unittest.skipUnless(HAS_PYMUPDF, "pymupdf no instalado: no se puede extraer texto del PDF")
    def test_hostile_question_exports_and_survives_in_the_pdf(self):
        pdf = PDFExporter.generate_pdf(_data())
        self.assertTrue(pdf.startswith(b"%PDF-"))
        text = _pdf_text(pdf)
        self.assertIn(HOSTILE_QUESTION, text)
        self.assertIn("Notas Ejecutivas", text)
        self.assertIn("nota <img src=y> & z", text)
        self.assertIn("<b>MALAS</b> & bajas", text)
        self.assertIn("Impacto <b>fuerte</b> & medible", text)
        self.assertIn(HOSTILE_SQL, text)

    @unittest.skipUnless(HAS_PYMUPDF, "pymupdf no instalado: no se puede extraer texto del PDF")
    def test_chart_that_cannot_be_decoded_is_declared_and_renumbered(self):
        data = _data(chart_image_base64="esto-no-es-una-imagen")
        text = _pdf_text(PDFExporter.generate_pdf(data))
        self.assertIn("Imagen no disponible", text)
        numbers = _section_numbers(text)
        self.assertEqual(numbers, list(range(1, len(numbers) + 1)), f"numeracion con hueco: {numbers}")

    @unittest.skipUnless(HAS_PYMUPDF, "pymupdf no instalado: no se puede extraer texto del PDF")
    def test_section_numbering_follows_what_was_rendered(self):
        valid_png = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        with_chart = _section_numbers(_pdf_text(PDFExporter.generate_pdf(_data(chart_image_base64=valid_png))))
        without_chart = _section_numbers(_pdf_text(PDFExporter.generate_pdf(_data())))
        self.assertEqual(with_chart, list(range(1, len(with_chart) + 1)))
        self.assertEqual(without_chart, list(range(1, len(without_chart) + 1)))
        # la figura existe, asi que no hay nota de imagen faltante y hay una seccion mas
        self.assertNotIn("Imagen no disponible", _pdf_text(PDFExporter.generate_pdf(_data(chart_image_base64=valid_png))))
        self.assertEqual(len(with_chart), len(without_chart) + 1)

    @unittest.skipUnless(HAS_PYMUPDF, "pymupdf no instalado: no se puede extraer texto del PDF")
    def test_seven_kpis_are_all_in_the_pdf(self):
        text = _pdf_text(PDFExporter.generate_pdf(_data()))
        for i in range(7):
            self.assertIn(f"KPI {i}", text)
        self.assertNotIn("trunc", text.lower())

    # --- BUG 1: el sanitizado tiene que alcanzar todas las celdas ---------

    def test_no_cell_is_stored_as_a_formula(self):
        content = ExcelExporter.generate_excel(_data())
        offenders = [
            (sheet, coord, cell.value)
            for sheet, coord, cell in _cells(content)
            if cell.data_type == "f"
        ]
        self.assertEqual(offenders, [])

    def test_hostile_values_are_neutralised_in_both_sheets(self):
        data = _data(question=FORMULA_QUESTION)
        wb = load_workbook(io.BytesIO(ExcelExporter.generate_excel(data)))
        summary, raw = wb["Resumen Ejecutivo"], wb["Datos"]
        # la pregunta viene del chat del usuario: no puede quedar viva en la hoja Resumen
        self.assertEqual(summary["B4"].value, "'" + FORMULA_QUESTION)
        self.assertEqual(raw["A2"].value, "'" + HOSTILE_SQL)
        trace_sql = [c.value for row in summary.iter_rows() for c in row if c.value == "'" + HOSTILE_SQL]
        self.assertTrue(trace_sql, "el SQL del snapshot tambien tiene que quedar neutralizado")
        # y el texto que no es formula se conserva tal cual (no se mutila el dato)
        self.assertEqual(load_workbook(io.BytesIO(ExcelExporter.generate_excel(_data())))["Datos"]["A3"].value, "normal")

    # --- BUG 4 y BUG 5: el audit dice lo que paso, y un log sin snapshot es 409 ---

    def test_export_without_raw_data_audits_the_original_row_count(self):
        entry = self._log(json.dumps(SNAPSHOT))

        response = self.client.post(
            "/api/v1/reports/export/excel",
            json={"audit_log_id": entry.id, "include_raw_data": False},
            headers=self.headers,
        )
        self.assertEqual(response.status_code, 200)

        export_log = self._last_export_log("EXPORTADO_EXCEL")
        self.assertEqual(export_log.rows_returned, 5000)

    def test_export_pdf_without_raw_data_audits_the_original_row_count(self):
        entry = self._log(json.dumps(SNAPSHOT))

        response = self.client.post(
            "/api/v1/reports/export/pdf",
            json={"audit_log_id": entry.id, "include_raw_data": False},
            headers=self.headers,
        )
        self.assertEqual(response.status_code, 200)

        export_log = self._last_export_log("EXPORTADO_PDF")
        self.assertEqual(export_log.rows_returned, 5000)

    def test_audit_row_count_falls_back_to_the_snapshot_rows(self):
        snapshot = dict(SNAPSHOT, traceability=None)
        entry = self._log(json.dumps(snapshot))

        response = self.client.post(
            "/api/v1/reports/export/excel",
            json={"audit_log_id": entry.id, "include_raw_data": False},
            headers=self.headers,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._last_export_log("EXPORTADO_EXCEL").rows_returned, 5)

    def test_audit_log_without_snapshot_returns_409_not_500(self):
        entry = self._log(None)
        for path in ("/api/v1/reports/export/pdf", "/api/v1/reports/export/excel"):
            with self.subTest(path=path):
                response = self.client.post(
                    path, json={"audit_log_id": entry.id}, headers=self.headers
                )
                self.assertEqual(response.status_code, 409)
                self.assertIn("no tiene resultado exportable", response.json().get("detail", ""))

    def test_corrupt_snapshot_returns_409_not_500(self):
        entry = self._log("{no soy json")
        for path in ("/api/v1/reports/export/pdf", "/api/v1/reports/export/excel"):
            with self.subTest(path=path):
                response = self.client.post(
                    path, json={"audit_log_id": entry.id}, headers=self.headers
                )
                self.assertEqual(response.status_code, 409)


if __name__ == "__main__":
    unittest.main()