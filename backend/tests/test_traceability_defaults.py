"""
Capa de origen del bug: el default de `TraceabilityAuditData.validation_status`.

Era `Optional[str] = "APROBADO"`. Un snapshot de auditoria guardado sin ese
campo (o con `null`) entra al export por `generator.py`, Pydantic le aplicaba el
default y el informe ejecutivo imprimia "Estado AST: APROBADO" sobre una consulta
que nadie valido. Los exporters ya saben decir "Sin registro de validacion"; con
este default nunca lo alcanzaban.

Regla del modulo: si no lo se, es None. Ningun default puede afirmar un resultado.
"""
import io
import json
import unittest

from openpyxl import load_workbook

from app.modules.admin_catalog.schemas import ReportExportData, TraceabilityAuditData
from app.modules.reports.excel_exporter import ExcelExporter
from app.modules.reports.pdf_exporter import PDFExporter


def _snapshot(traceability):
    """Snapshot tal como lo persiste chat_engine y lo relee el export."""
    return {
        "question": "Como van las ventas?",
        "data_columns": ["mes"],
        "data_rows": [{"mes": "enero"}],
        "traceability": traceability,
    }


class TestValidationStatusNeverDefaults(unittest.TestCase):
    def test_missing_validation_status_is_none_not_approved(self):
        # El schema no puede afirmar por su cuenta que la consulta fue validada.
        self.assertIsNone(TraceabilityAuditData().validation_status)
        self.assertIsNone(
            TraceabilityAuditData(sql_executed="SELECT 1", execution_time_ms=5, rows_returned=1).validation_status
        )

    def test_snapshot_without_validation_status_prints_no_approved(self):
        # El camino vivo: generator.py: json.loads(result_snapshot) -> ReportExportData(**dict).
        # Sin el campo, el default viejo lo rellenaba con "APROBADO" antes del export.
        snapshot = _snapshot({"sql_executed": "SELECT 1", "execution_time_ms": 5, "rows_returned": 1})
        self.assertNotIn("validation_status", snapshot["traceability"])

        data = ReportExportData(**json.loads(json.dumps(snapshot)))
        self.assertIsNone(data.traceability.validation_status)

        for content in (ExcelExporter.generate_excel(data), PDFExporter.generate_pdf(data)):
            self.assertNotIn(b"APROBADO", content)

    def test_snapshot_without_traceability_prints_no_approved(self):
        data = ReportExportData(**json.loads(json.dumps(_snapshot(None))))
        for content in (ExcelExporter.generate_excel(data), PDFExporter.generate_pdf(data)):
            self.assertNotIn(b"APROBADO", content)

    def test_missing_status_states_it_was_not_validated(self):
        data = ReportExportData(**json.loads(json.dumps(_snapshot({"sql_executed": "SELECT 1"}))))
        summary = load_workbook(io.BytesIO(ExcelExporter.generate_excel(data)))["Resumen Ejecutivo"]
        trace_line = next(
            c.value for row in summary.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("Estado AST:")
        )
        self.assertIn("Sin registro de validación", trace_line)

    def test_explicit_rejected_status_is_preserved(self):
        # El fix no puede pisar el dato real que el audit si trajo.
        for status in ("RECHAZADO", "RECHAZADO_RBAC", "ERROR_EJECUCION", "FUERA_DE_ALCANCE", "APROBADO_CONVERSACIONAL"):
            with self.subTest(status=status):
                data = ReportExportData(**json.loads(json.dumps(_snapshot({"sql_executed": "SELECT 1", "validation_status": status}))))
                self.assertEqual(data.traceability.validation_status, status)

                summary = load_workbook(io.BytesIO(ExcelExporter.generate_excel(data)))["Resumen Ejecutivo"]
                trace_line = next(
                    c.value for row in summary.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("Estado AST:")
                )
                self.assertIn(status, trace_line)
                self.assertNotIn("Sin registro de validación", trace_line)


if __name__ == "__main__":
    unittest.main()