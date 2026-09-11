import os
import re
import csv
import datetime
from typing import List, Optional, Any, Dict
from sqlalchemy import text, inspect
from .base import sanitize_identifier

def infer_postgres_type(values: List[Any]) -> str:
    """
    Infers the most appropriate PostgreSQL column type from a sample of non-null values.
    Returns: BIGINT, NUMERIC(14, 2), BOOLEAN, DATE, TIMESTAMP, or TEXT.
    """
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return "TEXT"

    # Check Boolean
    bool_strings = {"true", "false", "t", "f", "1", "0", "si", "no", "yes"}
    is_bool = True
    for v in non_empty:
        if isinstance(v, bool):
            continue
        if str(v).strip().lower() not in bool_strings:
            is_bool = False
            break
    if is_bool and len(non_empty) > 0 and all(isinstance(v, bool) or str(v).strip().lower() in ("true", "false", "si", "no") for v in non_empty):
        return "BOOLEAN"

    # Check Integer
    is_int = True
    for v in non_empty:
        s = str(v).strip().replace("$", "").replace("€", "").replace("%", "").replace(" ", "")
        try:
            int(s)
        except ValueError:
            is_int = False
            break
    if is_int:
        return "BIGINT"

    # Check Float / Numeric
    is_float = True
    for v in non_empty:
        s = str(v).strip().replace("$", "").replace("€", "").replace("%", "").replace(" ", "")
        if "," in s and "." not in s:
            s = s.replace(",", ".")
        try:
            float(s)
        except ValueError:
            is_float = False
            break
    if is_float:
        return "NUMERIC(14, 2)"

    # Check Date / Timestamp
    is_date = True
    is_timestamp = True
    for v in non_empty:
        if isinstance(v, datetime.date) and not isinstance(v, datetime.datetime):
            continue
        if isinstance(v, datetime.datetime):
            is_date = False
            continue
        s = str(v).strip()
        if len(s) == 10 and ((s[4] == '-' and s[7] == '-') or (s[4] == '/' and s[7] == '/')):
            try:
                datetime.date.fromisoformat(s.replace('/', '-'))
                continue
            except ValueError:
                pass
        is_date = False
        try:
            datetime.datetime.fromisoformat(s.replace('Z', '+00:00'))
        except ValueError:
            is_timestamp = False
            break
    if is_date:
        return "DATE"
    if is_timestamp:
        return "TIMESTAMP"

    return "TEXT"

def cast_postgres_val(val: Any, pg_type: str) -> Any:
    """
    Cleans and casts raw values to proper Python types suitable for PostgreSQL parameter binding.
    """
    if val is None:
        return None
    s = str(val).strip()
    if s == "" or s.lower() in ("nan", "none", "null", "n/a"):
        return None

    if pg_type == "BIGINT":
        s_num = s.replace("$", "").replace("€", "").replace("%", "").replace(" ", "")
        try:
            return int(float(s_num))
        except (ValueError, OverflowError):
            return None

    if pg_type == "NUMERIC(14, 2)":
        s_num = s.replace("$", "").replace("€", "").replace("%", "").replace(" ", "")
        if "," in s_num and "." not in s_num:
            s_num = s_num.replace(",", ".")
        try:
            return round(float(s_num), 2)
        except (ValueError, OverflowError):
            return None

    if pg_type == "BOOLEAN":
        if isinstance(val, bool):
            return val
        return s.lower() in ("true", "t", "1", "si", "yes")

    if pg_type == "DATE":
        if isinstance(val, datetime.date):
            return val
        try:
            return datetime.date.fromisoformat(s[:10].replace('/', '-'))
        except ValueError:
            return s[:10]

    if pg_type == "TIMESTAMP":
        if isinstance(val, datetime.datetime):
            return val
        try:
            return datetime.datetime.fromisoformat(s.replace('Z', ''))
        except ValueError:
            return s

    return s

def get_unique_table_name(engine, base_name: str) -> str:
    """
    Ensures table name does not collide with existing tables in PostgreSQL 'public' schema.
    Appends an incremental suffix (_1, _2, ...) if collision occurs.
    """
    inspector = inspect(engine)
    existing_tables = set(inspector.get_table_names(schema="public"))
    clean_name = sanitize_identifier(base_name, fallback_prefix="tbl")
    target_name = clean_name
    counter = 1
    while target_name in existing_tables:
        target_name = f"{clean_name}_{counter}"
        counter += 1
    return target_name

def create_table_and_insert(
    engine,
    table_name: str,
    headers: List[str],
    col_types: List[str],
    rows: List[List[Any]]
):
    """
    Creates table in PostgreSQL schema 'public' and inserts rows using chunked batch execution.
    """
    cols_def = ", ".join([f'"{h}" {t}' for h, t in zip(headers, col_types)])
    create_sql = f'CREATE TABLE IF NOT EXISTS "{table_name}" ({cols_def});'

    col_names = ", ".join([f'"{h}"' for h in headers])
    param_names = ", ".join([f':p_{idx}' for idx in range(len(headers))])
    insert_sql = text(f'INSERT INTO "{table_name}" ({col_names}) VALUES ({param_names});')

    prepared_data = []
    for r in rows:
        row_dict = {}
        for idx, (h, t) in enumerate(zip(headers, col_types)):
            val = r[idx] if idx < len(r) else None
            row_dict[f'p_{idx}'] = cast_postgres_val(val, t)
        prepared_data.append(row_dict)

    with engine.begin() as conn:
        conn.execute(text(f'DROP TABLE IF EXISTS "{table_name}";'))
        conn.execute(text(create_sql))
        chunk_size = 500
        for i in range(0, len(prepared_data), chunk_size):
            chunk = prepared_data[i:i + chunk_size]
            conn.execute(insert_sql, chunk)

def import_csv_to_postgres(csv_path: str, target_engine, table_name: Optional[str] = None) -> List[str]:
    """
    Imports a CSV/TSV file directly into PostgreSQL with automatic delimiter sniffing and type inference.
    """
    encodings = ["utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1"]
    content = None
    used_encoding = "utf-8"

    for enc in encodings:
        try:
            with open(csv_path, "r", encoding=enc) as f:
                content = f.read(65536)
                used_encoding = enc
                break
        except UnicodeDecodeError:
            continue

    if content is None:
        raise ValueError("No se pudo decodificar el archivo CSV con los juegos de caracteres estándar.")

    sample = content[:4096]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
        delimiter = dialect.delimiter
    except Exception:
        counts = {d: sample.count(d) for d in [",", ";", "\t", "|"]}
        delimiter = max(counts, key=counts.get) if any(counts.values()) else ","

    with open(csv_path, "r", encoding=used_encoding, errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        raw_headers = None
        data_rows = []
        for row in reader:
            if not row or all(str(cell).strip() == "" for cell in row):
                continue
            if raw_headers is None:
                raw_headers = row
            else:
                data_rows.append(row)

    if not raw_headers:
        raise ValueError("El archivo CSV está vacío o no contiene encabezados legibles.")

    headers = []
    seen_cols = set()
    for idx, h in enumerate(raw_headers):
        clean_col = sanitize_identifier(h, fallback_prefix=f"col_{idx+1}")
        col_final = clean_col
        c_count = 1
        while col_final in seen_cols:
            col_final = f"{clean_col}_{c_count}"
            c_count += 1
        seen_cols.add(col_final)
        headers.append(col_final)

    if not table_name:
        base_file = os.path.splitext(os.path.basename(csv_path))[0]
        base_file = re.sub(r'^raw_\d+_', '', base_file)
        table_name = sanitize_identifier(base_file, fallback_prefix="tabla_csv")

    col_types = []
    for c_idx in range(len(headers)):
        col_samples = [r[c_idx] for r in data_rows[:100] if c_idx < len(r)]
        col_types.append(infer_postgres_type(col_samples))

    pg_table_name = get_unique_table_name(target_engine, table_name)
    create_table_and_insert(target_engine, pg_table_name, headers, col_types, data_rows)
    return [pg_table_name]

def import_excel_to_postgres(excel_path: str, target_engine) -> List[str]:
    """
    Imports an Excel workbook (.xlsx, .xlsm, .xltx) into PostgreSQL.
    Creates a dedicated table for each non-empty worksheet.
    """
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    created_tables = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            raw_rows = list(ws.iter_rows(values_only=True))
            if not raw_rows:
                continue

            raw_headers = None
            data_rows = []
            for row in raw_rows:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                if raw_headers is None:
                    raw_headers = row
                else:
                    data_rows.append(row)

            if not raw_headers:
                continue

            headers = []
            seen_cols = set()
            for idx, h in enumerate(raw_headers):
                clean_col = sanitize_identifier(str(h) if h is not None else "", fallback_prefix=f"col_{idx+1}")
                col_final = clean_col
                c_count = 1
                while col_final in seen_cols:
                    col_final = f"{clean_col}_{c_count}"
                    c_count += 1
                seen_cols.add(col_final)
                headers.append(col_final)

            col_types = []
            for c_idx in range(len(headers)):
                col_samples = [r[c_idx] for r in data_rows[:100] if c_idx < len(r)]
                col_types.append(infer_postgres_type(col_samples))

            pg_table_name = get_unique_table_name(target_engine, sheet_name)
            create_table_and_insert(target_engine, pg_table_name, headers, col_types, data_rows)
            created_tables.append(pg_table_name)
    finally:
        wb.close()

    if not created_tables:
        raise ValueError("El archivo Excel no contiene hojas de cálculo con datos tabulares válidos.")
    return created_tables

def import_sqlite_to_postgres(sqlite_path: str, target_engine) -> List[str]:
    """
    Inspects an uploaded SQLite file and migrates all user tables into PostgreSQL.
    """
    import sqlite3
    conn = sqlite3.connect(sqlite_path)
    created_tables = []
    try:
        cur = conn.cursor()
        tables = [
            r[0] for r in cur.execute("SELECT name FROM sqlite_master WHERE type='table';").fetchall()
            if not r[0].startswith("sqlite_")
        ]
        if not tables:
            raise ValueError("El archivo SQLite subido no contiene tablas válidas.")

        for tbl in tables:
            pragma_cols = cur.execute(f'PRAGMA table_info("{tbl}");').fetchall()
            headers = [sanitize_identifier(col[1]) for col in pragma_cols]

            raw_rows = cur.execute(f'SELECT * FROM "{tbl}";').fetchall()
            col_types = []
            for idx in range(len(headers)):
                samples = [r[idx] for r in raw_rows[:100] if idx < len(r)]
                col_types.append(infer_postgres_type(samples))

            pg_table_name = get_unique_table_name(target_engine, tbl)
            create_table_and_insert(target_engine, pg_table_name, headers, col_types, raw_rows)
            created_tables.append(pg_table_name)
    finally:
        conn.close()

    return created_tables

def import_sql_script_to_postgres(sql_path: str, target_engine) -> List[str]:
    """
    Executes a raw SQL script against target PostgreSQL database.
    """
    with open(sql_path, "r", encoding="utf-8", errors="ignore") as sql_file:
        sql_script = sql_file.read()

    before_tables = set(inspect(target_engine).get_table_names(schema="public"))
    statements = [stmt.strip() for stmt in sql_script.split(";") if stmt.strip()]
    with target_engine.begin() as conn:
        for stmt in statements:
            try:
                conn.execute(text(stmt))
            except Exception:
                pass
    after_tables = set(inspect(target_engine).get_table_names(schema="public"))
    new_tables = list(after_tables - before_tables)
    return new_tables if new_tables else list(after_tables)
